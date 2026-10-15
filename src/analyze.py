"""Reproducible economic forecasts and an explicitly limited NLP pilot.

No API keys, LLM calls, fabricated human labels, or future-vintage predictors.
"""
from pathlib import Path
import sys,io,zipfile,json,re,math,hashlib,collections,csv
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'.build'/'pythonlibs'))
import numpy as np
import pandas as pd
import xlrd
from openpyxl import load_workbook

RAW=ROOT/'data/raw';OUT=ROOT/'data/processed';RES=ROOT/'results'
for p in [OUT,RES,ROOT/'annotations',ROOT/'figures',ROOT/'deliverables']:p.mkdir(parents=True,exist_ok=True)
SEED=20261014

def save_json(path,obj):
    def conv(x):
        if isinstance(x,(np.integer,np.floating)):return x.item()
        if isinstance(x,(pd.Timestamp,Path)):return str(x)
        raise TypeError(type(x))
    Path(path).write_text(json.dumps(obj,indent=2,default=conv,allow_nan=False),encoding='utf-8')

def load_snapshot(origin):
    z=zipfile.ZipFile(RAW/'vintages'/f'{origin}.zip');frames=[]
    for n in z.namelist():
        if n.endswith('.csv'):
            d=pd.read_csv(io.BytesIO(z.read(n)),index_col=0,parse_dates=True)
            d.columns=[c.rsplit('_',1)[0] for c in d.columns]
            frames.append(d)
    return pd.concat(frames,axis=1,sort=True)

def ratio_at_last(s,months):
    s=s.dropna();d=s.index[-1];prior=d-pd.DateOffset(months=months)
    return float(100*(s.iloc[-1]/s.get(prior,np.nan)-1))

def build_features():
    rows=[];snapshots={}
    manifest=json.loads((RAW/'vintage_manifest.json').read_text())
    for m in manifest:
        if not m.get('headers_verified'):raise ValueError('Unverified vintage '+m['origin'])
        origin=m['origin'];d=load_snapshot(origin);snapshots[origin]=d
        u=d.UNRATE.dropna();p=d.PCEPI.dropna();g=d.GDPC1.dropna();inc=d.DSPIC96.dropna()
        sent=d.UMCSENT.loc[d.index<pd.Timestamp(origin)].dropna()
        rows.append({'month':origin,'vintage_date':m['vintage_date'],'unemployment':u.iloc[-1],
          'inflation':ratio_at_last(p,12),'gdp_growth':100*((g.iloc[-1]/g.iloc[-2])**4-1),
          'income_growth':ratio_at_last(inc,12),'price_3y':ratio_at_last(p,36),'sentiment_last':sent.iloc[-1],
          'unemployment_observation':u.index[-1].date().isoformat(),'pce_observation':p.index[-1].date().isoformat(),
          'gdp_observation':g.index[-1].date().isoformat(),'sentiment_observation':sent.index[-1].date().isoformat(),
          'season_sin':math.sin(2*math.pi*pd.Timestamp(origin).month/12),'season_cos':math.cos(2*math.pi*pd.Timestamp(origin).month/12),
          'web_transition':min(1,max(0,(pd.Timestamp(origin).year*12+pd.Timestamp(origin).month-(2024*12+3))/4))})
    f=pd.DataFrame(rows).set_index('month');f.to_csv(OUT/'forecast_features.csv')
    # Targets for scoring are the current official historical series, supplemented
    # by September final and October preliminary figures verified at Michigan.
    target=pd.read_csv(RAW/'UMCSENT.csv',index_col=0,parse_dates=True).UMCSENT
    target.loc[pd.Timestamp('2026-09-01')]=48.1;target.loc[pd.Timestamp('2026-10-01')]=46.3
    target.sort_index().to_csv(OUT/'sentiment_targets.csv')
    save_json(OUT/'current_survey.json',{'source':'https://www.sca.isr.umich.edu/','retrieved_date':'2026-10-14','2026-10':{'status':'preliminary','sentiment':46.3,'current_conditions':44.7,'expectations':47.3,'release_date':'2026-10-09'},'2026-09':{'status':'final','sentiment':48.1,'current_conditions':50.9,'expectations':46.3},'2025-10':{'status':'final','sentiment':53.6,'current_conditions':58.6,'expectations':50.3}})
    return f,target,snapshots

def audit_supplied():
    w=load_workbook(ROOT/'Perceptual_Bias_Data.xlsx',data_only=True);s=w.active;results=[]
    for dc,vc,sid in [(1,2,'UNRATE'),(4,5,'UMCSENT'),(8,9,'PCE_YOY'),(11,12,'A191RL1Q225SBEA')]:
        source=pd.Series({pd.Timestamp(s.cell(r,dc).value).to_period('M').to_timestamp():s.cell(r,vc).value for r in range(2,s.max_row+1) if s.cell(r,dc).value is not None})
        if sid=='PCE_YOY':
            q=pd.read_csv(RAW/'PCEPI.csv',index_col=0,parse_dates=True).PCEPI;official=100*(q/q.shift(12)-1)
        else:official=pd.read_csv(RAW/(sid+'.csv'),index_col=0,parse_dates=True).iloc[:,0]
        if sid=='A191RL1Q225SBEA':source.index=source.index.to_period('Q').to_timestamp()
        joined=pd.concat([source.rename('supplied'),official.rename('official')],axis=1,sort=True).dropna()
        joined['difference']=joined.supplied-joined.official
        joined.to_csv(RES/f'audit_{sid}.csv')
        results.append({'series':sid,'supplied_n':len(source),'matched_n':len(joined),'mismatches_over_0_05':int((joined.difference.abs()>.05).sum()),'max_absolute_difference':float(joined.difference.abs().max()),'latest_matched':str(joined.index[-1].date()),'latest_supplied':float(joined.supplied.iloc[-1]),'latest_official':float(joined.official.iloc[-1])})
    save_json(RES/'data_audit.json',{'series':results,'excluded_future_gdp':{'period':'2026Q3','supplied_value':3.7,'reason':'Advance GDP not scheduled for publication until 2026-10-29'},'unemployment_missing':'2025-10: CPS not collected during federal shutdown; latest available observation used with its date retained','source_note':'Official FRED/ALFRED values replace the supplied workbook; matched differences may include revisions.'})

def ols(X,y):
    X=np.asarray(X,float);y=np.asarray(y,float);mu=X.mean(0);sd=X.std(0);sd[sd<1e-9]=1
    z=np.column_stack([np.ones(len(X)),(X-mu)/sd]);beta=np.linalg.lstsq(z,y,rcond=None)[0]
    return {'beta':beta,'mu':mu,'sd':sd,'residual':y-z@beta,'z':z}

def predict(fit,X):
    X=np.atleast_2d(np.asarray(X,float));z=np.column_stack([np.ones(len(X)),(X-fit['mu'])/fit['sd']]);return z@fit['beta']

def bootstrap_mean_ci(values,block=6,draws=2000):
    a=np.asarray(values,float);rng=np.random.default_rng(SEED);n=len(a);means=[]
    for _ in range(draws):
        idx=np.concatenate([(start+np.arange(block))%n for start in rng.integers(0,n,size=math.ceil(n/block))])[:n]
        means.append(a[idx].mean())
    return [float(v) for v in np.quantile(means,[.025,.975])]

def forecast_models(f,target,snapshots):
    specs={'Persistence':['sentiment_last'],
      'Economic':['unemployment','inflation','gdp_growth'],
      'Economic plus persistence':['unemployment','inflation','gdp_growth','sentiment_last'],
      'Economic plus household controls':['unemployment','inflation','gdp_growth','sentiment_last','income_growth','price_3y'],
      'Economic plus seasonal and method controls':['unemployment','inflation','gdp_growth','sentiment_last','income_growth','price_3y','season_sin','season_cos','web_transition']}
    predictions=[];coefs=[]
    for model,cols in specs.items():
        for origin,row in f.iterrows():
            if origin<'2018-01-01':continue
            # Fit only targets actually present in the forecast-date vintage.
            y_available=snapshots[origin].UMCSENT.dropna();y_available.index=y_available.index.strftime('%Y-%m-%d')
            train=f.loc[f.index<origin,cols].join(y_available.rename('target'),how='inner').dropna()
            if len(train)<60:continue
            if model=='Persistence':
                pred=float(row.sentiment_last)
            else:
                fit=ols(train[cols],train.target);pred=float(predict(fit,row[cols])[0])
                for c,b,sd in zip(cols,fit['beta'][1:],fit['sd']):coefs.append({'model':model,'month':origin,'feature':c,'coefficient':float(b/sd)})
            actual=float(target.get(pd.Timestamp(origin),np.nan));error=actual-pred
            # Prediction interval uses prior genuine out-of-sample errors only.
            calibration_targets=snapshots[origin].UMCSENT.dropna()
            prev=[float(calibration_targets.loc[pd.Timestamp(p['month'])])-p['prediction'] for p in predictions if p['model']==model and p['month']<origin and pd.Timestamp(p['month']) in calibration_targets.index]
            if len(prev)>=24:
                cal=np.asarray(prev[-60:]);radius=float(np.quantile(np.abs(cal),min(1,math.ceil((len(cal)+1)*.95)/len(cal)),method='higher'))
                lo,hi=pred-radius,pred+radius;method='95% trailing absolute forecast-error interval'
            else:lo=hi=None;method='insufficient prior forecast errors'
            predictions.append({'model':model,'month':origin,'prediction':pred,'actual':actual,'error':error,'lower':lo,'upper':hi,'interval_method':method,'training_n':len(train),'target_status':'preliminary' if origin=='2026-10-01' else 'final','vintage_date':row.vintage_date})
    p=pd.DataFrame(predictions);p.to_csv(RES/'forecasts.csv',index=False);pd.DataFrame(coefs).to_csv(RES/'coefficient_paths.csv',index=False)
    metrics=[]
    for label,start,end in [('All scored','2018-01-01','2026-09-01'),('Pre pandemic','2018-01-01','2019-12-01'),('Pandemic','2020-01-01','2021-12-01'),('Inflation and recovery','2022-01-01','2026-09-01')]:
        for model,g in p[(p.month>=start)&(p.month<=end)].groupby('model',sort=False):
            z=g.dropna(subset=['lower']);e=g.error.to_numpy();den=np.sum((g.actual-g.actual.mean())**2)
            metrics.append({'window':label,'model':model,'n':len(g),'RMSE':float(np.sqrt(np.mean(e**2))),'MAE':float(np.abs(e).mean()),'mean_error':float(e.mean()),'OOS_R2_vs_window_mean':float(1-np.sum(e**2)/den),'interval_n':len(z),'interval_coverage':float(((z.actual>=z.lower)&(z.actual<=z.upper)).mean()) if len(z) else None})
    save_json(RES/'forecast_metrics.json',metrics)
    # Fixed, prespecified comparison to persistence. Paired circular-block bootstrap.
    paired=p[p.month<='2026-09-01'].pivot(index='month',columns='model',values='error');comparisons=[]
    for model in specs:
        if model=='Persistence':continue
        loss=paired[model].to_numpy()**2-paired.Persistence.to_numpy()**2
        comparisons.append({'model':model,'mean_squared_error_minus_persistence':float(loss.mean()),'block_bootstrap_95_ci':bootstrap_mean_ci(loss)})
    save_json(RES/'model_comparisons.json',comparisons)
    # Sensitivity: fixed 2010–2019 training information, predicting later years.
    sensitivity=[];origin='2020-01-01';available=snapshots[origin].UMCSENT.dropna();available.index=available.index.strftime('%Y-%m-%d')
    for model in ['Economic','Economic plus persistence','Economic plus household controls']:
        cols=specs[model];tr=f.loc[f.index<origin,cols].join(available.rename('target')).dropna();fit=ols(tr[cols],tr.target)
        te=f.loc[(f.index>=origin)&(f.index<='2026-09-01')];actual=target.reindex(pd.to_datetime(te.index)).to_numpy();e=actual-predict(fit,te[cols]);sensitivity.append({'model':model,'training':'fixed 2010–2019 vintage available by 2019-12-31','test_n':len(e),'RMSE':float(np.sqrt(np.mean(e**2))),'MAE':float(np.abs(e).mean())})
    save_json(RES/'fixed_training_sensitivity.json',sensitivity)
    # HAC standard errors are descriptive, not causal or a model-selection rule.
    cols=specs['Economic'];origin=f.index[-1];available=snapshots[origin].UMCSENT.dropna();available.index=available.index.strftime('%Y-%m-%d');tr=f[cols].join(available.rename('target')).dropna();fit=ols(tr[cols],tr.target)
    Z=fit['z'];e=fit['residual'];bread=np.linalg.pinv(Z.T@Z);v=Z*e[:,None];meat=v.T@v
    for lag in range(1,7):
        s=v[lag:].T@v[:-lag];meat+=(1-lag/7)*(s+s.T)
    cov=bread@meat@bread*len(Z)/(len(Z)-Z.shape[1]);se=np.sqrt(np.maximum(0,np.diag(cov)))[1:]/fit['sd'];coef=fit['beta'][1:]/fit['sd']
    save_json(RES/'economic_coefficients_hac.json',[{'feature':c,'coefficient':float(b),'HAC_SE_6_lags':float(s),'normal_95_lower':float(b-1.96*s),'normal_95_upper':float(b+1.96*s)} for c,b,s in zip(cols,coef,se)])
    save_json(RES/'model_specifications.json',{'forecast_origin':'first day of target month','information_cutoff':'previous calendar day end of day','economic_vintages':'verified ALFRED snapshot for each origin','estimation_start':'2010-01','evaluation_start':'2018-01','final_scored_month':'2026-09','monitoring_only_preliminary_month':'2026-10','models':specs,'media_models':{'Economic plus media':'not estimable: fewer than 24 historical monthly news observations meeting >=10 headlines per month','Economic plus media plus controls':'not estimable for the same reason'},'uncertainty':'rolling empirical intervals do not assume independent errors; nominal 95% coverage is not guaranteed under regime shifts; block-bootstrap comparisons use length 6 and 2000 replicates','training_labels':'latest available sentiment values in the origin vintage','scoring_labels':'latest final values as of cutoff, except October preliminary reported separately','interval_calibration':'prior forecasts scored against outcomes actually available in the current origin vintage; trailing 60 errors, minimum 24','political_controls':'not included as predictors: party-average sentiment is a survey outcome, not independent exposure or an exogenous control'})
    return p,metrics

NEG=set('crisis recession decline declines declining fell fall falls falling loss losses weak weakness worry worries worried fear fears pessimism pessimistic slump slumps struggling struggle layoffs layoffs layoff expensive unaffordable burden burdens squeeze squeezed soaring surge surges surged inflation uncertainty risk risks cut cuts downturn worsening worsen worse hardship shortages collapse disappointing slows slowed slowdown'.split())
POS=set('growth growing grow recovery recover recovered rebound rebounds gains gain gained strong strength improves improve improved improving optimism optimistic relief resilient resilience boom booming hiring eased easing eases cool cooling cooled lower increase increases rose rise rises rising'.split())
STOP=set('the a an and or of to in for on at with from by is are was were be as this that it its us u s after before says said has have had will would could may can about into than more'.split())

def tokens(text):return [t for t in re.findall(r'[a-z]+',text.lower()) if t not in STOP and len(t)>1]
def dictionary(text):
    t=tokens(text);score=sum(w in POS for w in t)-sum(w in NEG for w in t)
    return 'positive' if score>0 else 'negative' if score<0 else 'neutral'

def fit_nb(rows):
    vocab={t for r in rows for t in tokens(r['text'])};vocab={t:i for i,t in enumerate(sorted(vocab))};labels=['negative','neutral','positive'];counts=np.ones((3,len(vocab)));n=np.zeros(3)
    for r in rows:
        k=labels.index(r['label']);n[k]+=1
        for t in tokens(r['text']):counts[k,vocab[t]]+=1
    return {'vocab':vocab,'labels':labels,'logprob':np.log(counts/counts.sum(1)[:,None]),'logprior':np.log(n/n.sum())}

def nb_predict(model,text):
    score=model['logprior'].copy()
    for t in tokens(text):
        if t in model['vocab']:score+=model['logprob'][:,model['vocab'][t]]
    exp=np.exp(score-score.max());prob=exp/exp.sum();k=int(prob.argmax());return model['labels'][k],float(prob[k])

def classification_metrics(actual,pred):
    labels=['negative','neutral','positive'];cm=np.zeros((3,3),int)
    for a,b in zip(actual,pred):cm[labels.index(a),labels.index(b)]+=1
    classes=[]
    for k,l in enumerate(labels):
        precision=float(cm[k,k]/max(1,cm[:,k].sum()));recall=float(cm[k,k]/max(1,cm[k].sum()));f1=2*precision*recall/max(1e-12,precision+recall)
        classes.append({'label':l,'support':int(cm[k].sum()),'precision':precision,'recall':recall,'F1':f1})
    return {'accuracy':float(np.trace(cm)/max(1,cm.sum())),'macro_F1':float(np.mean([x['F1'] for x in classes])),'classes':classes,'confusion_matrix':cm.tolist(),'class_order':labels}

def text_measures(text):
    low=text.lower();flags={
      'salience_layoffs':r'layoff|job cuts|jobs cut|lost jobs|job losses',
      'salience_hiring':r'hir(?:e|ing)|jobs added|job gains|payroll',
      'salience_prices':r'inflation|prices|cost of living|grocer|rent|affordab',
      'salience_growth':r'\bgdp\b|growth|economy grows|economic expansion',
      'salience_recession':r'recession|downturn|economic contraction',
      'frame_price_level':r'cost of living|high prices|higher prices|expensive|affordab|grocery|groceries|rent',
      'frame_inflation_rate':r'inflation|disinflation|price growth|pce|cpi',
      'frame_household':r'household|famil|consumer|grocery|groceries|rent|paycheck|wallet|cost of living',
      'frame_aggregate':r'\bgdp\b|unemployment|payroll|national economy|us economy|u\.s\. economy',
      'frame_forecast':r'expect|forecast|outlook|will |could |may |risk|ahead',
      'political_attribution_mention':r'trump|biden|republican|democrat|white house|congress|federal reserve|\bfed\b',
      'emotional_language':r'crisis|disaster|panic|terrifying|devastat|booming|surging|soaring|collapse|crushing'}
    out={k:int(bool(re.search(v,low))) for k,v in flags.items()}
    out['numeric_claim_present']=int(bool(re.search(r'\d+(?:\.\d+)?\s*(?:%|percent|million|billion|jobs)',low)))
    # A matching number alone cannot establish a factual match. Human/source review is required.
    out['factual_calibration']='needs_source_review' if out['numeric_claim_present'] else 'no_numeric_claim_detected'
    return out

def calibrate_numeric_claim(row):
    text=row['text'].lower();date=row['published_date']
    percentages=[float(x) for x in re.findall(r'(\d+(?:\.\d+)?)\s*(?:%|percent)',text)]
    result={'calibration_benchmark':None,'calibration_claim':None,'calibration_official_value':None,'calibration_source':None}
    if not row['numeric_claim_present']:return result
    # Narrow, explicit claims only. Never compare CPI with PCE, or a forecast with a release.
    if row['frame_forecast']:return result
    if 'unemployment' in text and ('september' in text or 'sept.' in text) and date>='2026-10-02' and len(percentages)==1:
        official=4.2;source='https://www.bls.gov/news.release/archives/empsit_10022026.htm';benchmark='September 2026 U-3 released 2026-10-02'
    elif ('pce' in text or 'personal consumption expenditures' in text) and 'august' in text and date>='2026-09-30' and len(percentages)==1 and not re.search(r'core|monthly|month.on.month',text):
        official=3.4;source='https://www.bea.gov/news/2026/personal-income-and-outlays-august-2026';benchmark='August 2026 headline PCE YoY released 2026-09-30'
    else:return result
    result.update(calibration_benchmark=benchmark,calibration_claim=percentages[0],calibration_official_value=official,calibration_source=source)
    row['factual_calibration']='matches_narrow_benchmark' if abs(percentages[0]-official)<.051 else 'mismatches_narrow_benchmark'
    return result

def nlp_pilot():
    z=zipfile.ZipFile(RAW/'phrasebank.zip');body=z.read('FinancialPhraseBank-v1.0/Sentences_AllAgree.txt').decode('latin-1');raw=[];seen=set()
    for line in body.splitlines():
        if '@' not in line:continue
        text,label=line.rsplit('@',1);key=' '.join(tokens(text));
        if key in seen:continue
        seen.add(key);raw.append({'text':text,'label':label.strip()})
    rng=np.random.default_rng(SEED);train=[];test=[]
    for label in ['negative','neutral','positive']:
        group=[r for r in raw if r['label']==label];rng.shuffle(group);n=round(len(group)*.2);test+=group[:n];train+=group[n:]
    model=fit_nb(train);actual=[r['label'] for r in test];pred=[nb_predict(model,r['text'])[0] for r in test]
    benchmark={'dataset':'Financial PhraseBank unanimously annotated subset','license':'CC BY-NC-SA 3.0','source':'https://huggingface.co/datasets/takala/financial_phrasebank','seed':SEED,'deduplicated_n':len(raw),'train_n':len(train),'test_n':len(test),'split':'stratified 80/20 after normalized exact deduplication; near-duplicate and issuer overlap may remain','Naive Bayes':classification_metrics(actual,pred),'Dictionary':classification_metrics(actual,[dictionary(r['text']) for r in test]),'domain_limit':'company financial sentences differ from US macro headlines and Reddit titles; external benchmark does not validate the project corpus'}
    save_json(RES/'nlp_benchmark.json',benchmark)
    save_json(RES/'tone_model.json',{'algorithm':'multinomial_naive_bayes','version':'nb-phrasebank-v1','training_sha256':hashlib.sha256(body.encode()).hexdigest(),'vocab':model['vocab'],'labels':model['labels'],'logprob':model['logprob'].tolist(),'logprior':model['logprior'].tolist(),'seed':SEED})
    corpus=json.loads((OUT/'corpus.json').read_text());focused=[];archive=[]
    for row in corpus:
        if row['published_date']<'2026-07-01':archive.append(row);continue
        row.update(text_measures(row['text']));row.update(calibrate_numeric_claim(row));row['tone'],row['model_probability']=nb_predict(model,row['text']);row['dictionary_tone']=dictionary(row['text']);row['tone_disagreement']=int(row['tone']!=row['dictionary_tone']);row['model_version']='nb-phrasebank-v1';focused.append(row)
    save_json(OUT/'corpus_archive_sparse.json',archive);save_json(OUT/'corpus_coded.json',focused)
    frame=pd.DataFrame(focused);frame.to_csv(OUT/'corpus_coded.csv',index=False)
    frame['month']=frame.published_date.str[:7];frame['tone_score']=frame.tone.map({'negative':-1,'neutral':0,'positive':1})
    fields=[c for c in frame if c.startswith(('salience_','frame_'))]+['tone_score','tone_disagreement','political_attribution_mention','numeric_claim_present','emotional_language']
    monthly=frame.groupby(['month','source_kind'])[fields].mean();monthly['n']=frame.groupby(['month','source_kind']).size();monthly.to_csv(RES/'media_monthly_pilot.csv')
    source=frame.groupby(['source_kind','source']).agg(n=('item_id','count'),negative_share=('tone',lambda x:float((x=='negative').mean())),disagreement=('tone_disagreement','mean'));source.to_csv(RES/'source_composition.csv')
    # Stratify the 300-item queue by source type, predicted tone and month.
    groups={k:g.to_dict('records') for k,g in frame.groupby(['source_kind','tone','month'])};sample=[]
    for g in groups.values():rng.shuffle(g)
    while len(sample)<min(300,len(frame)):
        for key in sorted(groups):
            if groups[key] and len(sample)<300:sample.append(groups[key].pop())
        if not any(groups.values()):break
    columns=['item_id','text','source_kind','published_date','tone','price_level_frame','inflation_rate_frame','household_frame','aggregate_frame','forecast_frame','numeric_claim','factual_calibration','political_credit_blame','evidence_excerpt','annotator_id']
    # Blind files intentionally omit source names, URLs and all model outputs.
    for annotator in ['A','B']:
        annotation_file=ROOT/'annotations'/f'annotator_{annotator}.csv'
        if annotation_file.exists() and pd.read_csv(annotation_file).tone.notna().any():
            continue  # Preserve actual human work across reruns.
        with annotation_file.open('w',encoding='utf-8-sig',newline='') as f:
            w=csv.DictWriter(f,fieldnames=columns);w.writeheader()
            for r in sample:w.writerow({c:r.get(c,'') if c in ['item_id','text','source_kind','published_date'] else '' for c in columns})
    pd.DataFrame(sample)[['item_id','source','url','tone','dictionary_tone','model_probability']].to_csv(ROOT/'annotations'/'audit_key.csv',index=False)
    # Name-swap invariance probes measure model sensitivity, not truth or fairness.
    changes=[]
    for r in focused:
        if re.search(r'\b(trump|biden|republicans?|democrats?)\b',r['text'],re.I):
            def swap(m):return {'trump':'Biden','biden':'Trump','republican':'Democrat','republicans':'Democrats','democrat':'Republican','democrats':'Republicans'}[m.group(0).lower()]
            altered=re.sub(r'\b(trump|biden|republicans?|democrats?)\b',swap,r['text'],flags=re.I);new,_=nb_predict(model,altered);changes.append({'item_id':r['item_id'],'original_tone':r['tone'],'swapped_tone':new,'changed':new!=r['tone']})
    pd.DataFrame(changes).to_csv(RES/'name_swap_audit.csv',index=False)
    summary={'n':len(focused),'by_source_kind':frame.source_kind.value_counts().to_dict(),'by_month':frame.month.value_counts().sort_index().to_dict(),'tone_counts':frame.tone.value_counts().to_dict(),'dictionary_disagreement':float(frame.tone_disagreement.mean()),'numeric_claims_detected':int(frame.numeric_claim_present.sum()),'factual_calibration_counts':frame.factual_calibration.value_counts().to_dict(),'annotation_queue_n':len(sample),'human_labels_completed':0,'human_agreement':None,'project_corpus_macro_F1':None,'name_swap_probes':len(changes),'name_swap_changed':sum(x['changed'] for x in changes),'media_prediction_status':'not estimable with four focused monthly bins and no historical coverage','source_weighting':'item weighted by default; source-equal sensitivity also computed','partial_month':'October stops at 2026-10-14; compare cautiously with full months'}
    save_json(RES/'corpus_summary.json',summary)
    # Equal outlet weight sensitivity within source type, avoiding domination by prolific publishers.
    equal=frame.groupby(['month','source_kind','source'])[fields].mean().groupby(['month','source_kind']).mean();equal.to_csv(RES/'media_source_equal_pilot.csv')
    return summary

def survey_components():
    results={}
    for name in ['1bh','5bh','2h','3h']:
        book=xlrd.open_workbook(RAW/f'michigan_{name}.xls');s=book.sheet_by_name('Data');rows=[]
        for r in range(2,s.nrows):
            vals=s.row_values(r)
            if not isinstance(vals[0],(int,float)):continue
            date=xlrd.xldate_as_datetime(vals[0],book.datemode)
            if date.year<2010:continue
            rows.append([date.strftime('%Y-%m-%d')]+[v if isinstance(v,(float,int)) else None for v in vals[2:]])
        names=s.row_values(0)[2:];d=pd.DataFrame(rows,columns=['month']+names);d.to_csv(OUT/f'survey_{name}.csv',index=False)
        last=d.iloc[-1];results[name]={'latest_month':last['month'],'values':{n:float(last[n]) for n in names},'measurement_note':'party series labeled PLREPICS3MMA, PLDEMICS3MMA, PLINDICS3MMA; three-month moving averages' if name=='5bh' else 'descriptive series from Michigan August 2026 historical chart; no historical release vintages recovered'}
    save_json(RES/'survey_subgroups.json',results)

def charts(p,metrics):
    import matplotlib;matplotlib.use('Agg');import matplotlib.pyplot as plt
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,'axes.spines.right':False,'axes.grid':True,'grid.alpha':.2,'figure.dpi':140})
    colors={'actual':'#15324f','prediction':'#00867d','gap':'#b4533b'}
    for model,stem in [('Economic','economic_forecast'),('Economic plus persistence','dynamic_forecast')]:
        d=p[p.model==model].copy();d.month=pd.to_datetime(d.month)
        fig,ax=plt.subplots(figsize=(9,3.6));ax.plot(d.month,d.actual,label='Observed sentiment',color=colors['actual'],lw=1.8);ax.plot(d.month,d.prediction,label='Forecast',color=colors['prediction'],lw=1.5);ax.fill_between(d.month,d.lower,d.upper,color=colors['prediction'],alpha=.13,label='95% empirical interval');ax.set(ylabel='Michigan index points',title=model+' model using origin-date economic vintages');ax.legend(ncol=3,fontsize=8,loc='upper right');fig.tight_layout();fig.savefig(ROOT/'figures'/f'{stem}.png');plt.close(fig)
    d=p[p.model=='Economic'];fig,ax=plt.subplots(figsize=(9,3));ax.axhline(0,color='#777',lw=.8);ax.bar(pd.to_datetime(d.month),d.error,width=24,color=colors['gap']);ax.set(ylabel='Observed minus forecast',title='Sentiment gap under the economic-only specification');fig.tight_layout();fig.savefig(ROOT/'figures/gap.png');plt.close(fig)
    rows=[m for m in metrics if m['window']=='All scored'];fig,ax=plt.subplots(figsize=(9,3.4));labels=[r['model'].replace('Economic plus ','+ ') for r in rows];ax.barh(labels,[r['RMSE'] for r in rows],color=['#70879c','#b4533b','#00867d','#4576a0','#77835d']);ax.invert_yaxis();ax.set(xlabel='RMSE in Michigan index points',title='Chronological forecast accuracy January 2018 to September 2026');fig.tight_layout();fig.savefig(ROOT/'figures/model_comparison.png');plt.close(fig)
    d=pd.read_csv(RES/'media_monthly_pilot.csv');fig,ax=plt.subplots(figsize=(9,3.1));
    for kind,g in d.groupby('source_kind'):ax.plot(g.month,g.tone_score,marker='o',label=kind.title())
    ax.set(ylim=(-1,1),ylabel='Mean classified tone',title='Unvalidated headline and post-title tone in the pilot corpus');ax.legend(ncol=3);fig.tight_layout();fig.savefig(ROOT/'figures/media_pilot.png');plt.close(fig)

def main():
    audit_supplied();f,t,s=build_features();p,m=forecast_models(f,t,s);summary=nlp_pilot();survey_components();charts(p,m)
    print('Forecast rows:',len(p),'Corpus:',summary['n'])
    print(pd.DataFrame(m).query("window=='All scored'")[['model','n','RMSE','MAE','interval_coverage']].to_string(index=False))

if __name__=='__main__':main()
